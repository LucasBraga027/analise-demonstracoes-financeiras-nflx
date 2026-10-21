"""
Etapa 6 - Confere os valores que o Excel calculou na planilha contra um cálculo independente feito em pandas
sobre a mesma base. Rodar depois de recalcular/salvar a planilha no Excel (scripts/recalcular_excel.ps1).
"""
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

from sec_common import ROOT

base = pd.read_csv(os.path.join(ROOT, 'dados', 'base_trimestral_nflx.csv'), index_col=0)
wb = openpyxl.load_workbook(os.path.join(ROOT, 'Analise_Netflix_ECO057.xlsx'), data_only=True)


def linha_excel(aba, rotulo):
    ws = wb[aba]
    for row in ws.iter_rows(min_row=5):
        if row[0].value == rotulo:
            vals = [c.value for c in row[2:2 + len(base)]]
            return pd.Series([np.nan if (x is None or isinstance(x, str)) else x for x in vals], index=base.index, dtype=float)
    raise KeyError(f'{aba}: {rotulo}')


b = base
ltm = lambda s: s.rolling(4).sum()
esperado = {
    ('DRE', '(=) Lucro líquido'): b['lucro_liquido'],
    ('DRE', '(=) EBIT (EBIT / LAJIR)'.replace('(=) EBIT (EBIT / LAJIR)', '(=) Lucro operacional (EBIT / LAJIR)')): b['ebit'],
    ('BP', 'ATIVO CIRCULANTE'): b['ativo_circulante'],
    ('BP', 'PASSIVO TOTAL (exigível)'): b['passivo_total'],
    ('DFC', '(=) Fluxo de caixa livre simplificado (FCO − Capex)'): b['fco'] - b['capex'],
    ('AV', 'EBIT'): b['ebit'] / b['receita'],
    ('AH', 'Receita líquida'): (b['receita'] - b['receita'].shift(4)) / b['receita'].shift(4).abs(),
    ('Indicadores', 'ROE = ROA × MAF'): ltm(b['lucro_liquido']) / b['patrimonio_liquido'],
    ('Indicadores', 'NCG = ACO − PCO'): b['outros_ativos_circ'] - (b['passivo_circulante'] - b['divida_cp']),
    ('Indicadores', 'Saldo de tesouraria ST = CCL − NCG'): b['caixa'] + b['aplicacoes_cp'] - b['divida_cp'],
    ('Indicadores', 'Dívida líquida / EBITDA LTM'): (b['divida_cp'] + b['divida_lp'] - b['caixa'] - b['aplicacoes_cp'])
    / ltm(b['ebit'] + b['depreciacao_amort']),
    ('Indicadores', 'GAO = Δ%EBIT / Δ%Receita'): ((b['ebit'] / b['ebit'].shift(4) - 1) / (b['receita'] / b['receita'].shift(4) - 1)),
}
falhas = 0
for (aba, rot), esp in esperado.items():
    got = linha_excel(aba, rot)
    dif = (got - esp).abs().dropna()
    ok = len(dif) > 0 and dif.max() < 1e-6 * max(1, esp.abs().max())
    falhas += not ok
    print(f"{'OK ' if ok else 'ERRO'} {aba:12} {rot[:60]:60} n={len(dif):2d} maior dif={dif.max() if len(dif) else float('nan'):.2e}")
val = wb['Validacao']
for row in val.iter_rows(min_row=5, max_row=12):
    print(f'   Validacao: {row[0].value[:70]:70} {row[1].value} {row[2].value}')
    falhas += row[2].value != 'OK'
print('\nRESULTADO:', 'tudo confere' if falhas == 0 else f'{falhas} problema(s)')
sys.exit(1 if falhas else 0)
